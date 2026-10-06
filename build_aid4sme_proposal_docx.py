from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "AID4SME_Project_Proposal_EnerTwin_Ledger.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(20, 20, 20)
MUTED = RGBColor(90, 90, 90)
LIGHT_FILL = "F4F6F9"
GRID = "D9E2EC"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph_border_bottom(paragraph, color="D7DBE2", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, widths, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    set_run_font(run, size=9.5, color=INK)
            if header and r_idx == 0:
                set_cell_shading(cell, LIGHT_FILL)
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_run_font(run, size=9.5, color=DARK_BLUE, bold=True)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, color=DARK_BLUE, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.2, color=INK)
    doc.add_paragraph()


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Number"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AID4SME proposal draft | EnerTwin Ledger")
    set_run_font(r, size=8.5, color=MUTED)


def build_doc():
    doc = Document()
    configure_styles(doc)
    add_footer(doc.sections[0])

    # Cover/title block.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("AID4SME Open Call #2")
    set_run_font(r, size=12, color=MUTED, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("EnerTwin Ledger")
    set_run_font(r, size=26, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("AI-driven Digital Twin and Blockchain Trust Layer for Flexible Microgrid Energy Optimisation")
    set_run_font(r, size=13, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("Project proposal tailored to Challenge 2.3 - Energy system Digital Twin decision support tool")
    set_run_font(r, size=10.5, color=MUTED, bold=True)
    paragraph_border_bottom(p)

    meta = doc.add_table(rows=4, cols=4)
    rows = [
        ("Programme", "AID4SME Open Call #2", "Challenge", "2.3 Energy system Digital Twin"),
        ("Duration", "14 months", "Requested funding", "EUR 200,000"),
        ("Applicant type", "SME / start-up", "TRL target", "Validated TRL 6-7 prototype"),
        ("Core domains", "AI, energy optimisation", "Trust layer", "Blockchain verification"),
    ]
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = meta.cell(i, j)
            cell.text = value
            if j in (0, 2):
                set_cell_shading(cell, LIGHT_FILL)
    style_table(meta, [1800, 2880, 1800, 2880], header=False)
    doc.add_paragraph()

    add_callout(
        doc,
        "Core proposition",
        "EnerTwin Ledger combines a real-time microgrid Digital Twin, AI forecasting and predictive optimisation with a blockchain-based verification layer for trusted energy data, flexibility activation and ancillary service settlement.",
    )

    add_heading(doc, "1. Executive summary", 1)
    add_body(
        doc,
        "EnerTwin Ledger will develop an AI-driven decision support tool for microgrid operators that combines an energy system Digital Twin, predictive optimisation algorithms and a blockchain-based trust layer for data traceability, flexibility verification and ancillary service settlement. The solution will help operators optimise the sizing and operation of renewable generation, battery storage, thermal storage and flexible loads while reducing energy costs, increasing renewable self-consumption and enabling verifiable participation in grid flexibility markets.",
    )
    add_body(
        doc,
        "The project directly addresses AID4SME Challenge 2.3 by delivering optimisation algorithms for microgrid component sizing, control algorithms based on predicted heat and electricity needs, and validation in a digital environment using real-world data across multiple scenarios.",
    )

    add_heading(doc, "2. Alignment with AID4SME Challenge 2.3", 1)
    add_body(
        doc,
        "The proposal is positioned under Domain 2: Combined AI and Data solutions for creation of insights, specifically Challenge 2.3 - Energy system Digital Twin decision support tool. It responds to the expected contribution areas in the call by developing an optimisation algorithm for microgrid component sizing, developing control algorithms using predicted energy needs and production, testing the solution in a digital environment with real-world data, and demonstrating ancillary service scenarios together with thermal storage.",
    )
    add_bullet(doc, "AI and data focus: forecasting, anomaly-aware data pipelines, scenario simulation and model predictive control.")
    add_bullet(doc, "Digital Twin focus: coupled electrical and thermal energy balance with asset sizing, operational cost and pricing scenarios.")
    add_bullet(doc, "Blockchain focus: auditable records of metered data, flexibility events, baseline calculations and delivered ancillary service volumes.")
    add_bullet(doc, "Green Deal focus: higher renewable self-consumption, lower curtailment and more efficient use of distributed storage assets.")

    add_heading(doc, "3. Problem and need", 1)
    add_body(
        doc,
        "Microgrids that combine solar generation, batteries, thermal storage, heat pumps and industrial loads are increasingly difficult to size and operate. Operators need to balance cost, reliability, renewable energy use and grid service obligations under dynamic tariffs and variable demand profiles. Static sizing tools and rule-based control strategies are often insufficient because they do not learn from real operational data and cannot easily quantify the value of flexibility.",
    )
    add_body(
        doc,
        "A second barrier is trust. If microgrids are to participate in flexibility and ancillary service markets, energy data, asset status, baseline calculations, activation events and delivered service volumes must be auditable. The blockchain layer in this proposal is not used as a speculative trading feature; it is used as a practical verification mechanism for energy data provenance, event integrity and transparent settlement.",
    )

    add_heading(doc, "4. Proposed solution", 1)
    add_body(
        doc,
        "EnerTwin Ledger will be delivered as a modular software prototype composed of four integrated layers.",
    )
    components = doc.add_table(rows=1, cols=3)
    hdr = components.rows[0].cells
    hdr[0].text = "Layer"
    hdr[1].text = "Function"
    hdr[2].text = "AID4SME relevance"
    for layer, function, relevance in [
        ("Energy Digital Twin", "Simulates electrical and thermal energy flows, including PV, battery storage, thermal storage, heat demand, electricity demand, grid import/export and dynamic tariffs.", "Core digital environment for scenario testing and asset sizing."),
        ("AI forecasting", "Predicts electricity demand, heat demand, renewable generation and flexibility availability from real-world and historical data.", "Combined AI and data approach for insights and decision support."),
        ("Optimisation and control", "Performs microgrid component sizing and model predictive control for storage, load shifting and ancillary service activation.", "Direct response to sizing and control requirements in Challenge 2.3."),
        ("Blockchain trust layer", "Timestamped records for energy data, flexibility activation, delivery evidence and settlement-ready audit trails.", "Strengthens exploitation by adding verifiable performance and market trust."),
    ]:
        cells = components.add_row().cells
        cells[0].text = layer
        cells[1].text = function
        cells[2].text = relevance
    style_table(components, [1900, 4200, 3260])

    add_heading(doc, "5. Innovation", 1)
    add_body(
        doc,
        "The project is innovative because it combines three technologies that are often implemented separately: Digital Twin simulation, AI-driven optimisation and blockchain-based verification. The Digital Twin allows operators to test sizing and control scenarios before deployment. AI forecasting and optimisation transform the Digital Twin from a passive model into an operational decision-support tool. The blockchain layer adds a trustworthy record of data and flexibility events, which is essential for scaling from internal energy optimisation to external grid service participation.",
    )
    add_bullet(doc, "Novel integration of electrical and thermal energy storage optimisation with market-facing flexibility verification.")
    add_bullet(doc, "Scenario-based sizing that considers CAPEX, operational cost, renewable curtailment and ancillary service revenue potential.")
    add_bullet(doc, "Auditable flexibility delivery records suitable for microgrid operators, aggregators and energy communities.")

    add_heading(doc, "6. Technical objectives", 1)
    for objective in [
        "Develop a microgrid Digital Twin capable of simulating electrical and thermal energy balance under different asset sizing and pricing scenarios.",
        "Build AI forecasting models for electricity demand, heat demand and renewable generation.",
        "Develop optimisation algorithms for microgrid sizing and operational control.",
        "Implement a blockchain-based verification layer for energy data integrity and ancillary service delivery.",
        "Validate the prototype using real-world energy datasets and AID4SME digital playground scenarios.",
        "Prepare a business plan and exploitation roadmap for deployment with industrial microgrids, energy communities and flexibility aggregators.",
    ]:
        add_number(doc, objective)

    add_heading(doc, "7. Work plan and milestones", 1)
    plan = doc.add_table(rows=1, cols=4)
    for idx, text in enumerate(["Period", "Work package", "Main activities", "Output"]):
        plan.rows[0].cells[idx].text = text
    for row in [
        ("M1", "WP1 - Plan and alignment", "Refine use case with AID4SME mentors and playground owners; define data inputs, baseline scenarios, KPI values and IPR assumptions.", "Validated action plan and KPI baseline."),
        ("M2-M4", "WP2 - Digital Twin and data model", "Build microgrid simulation model; integrate electrical and thermal assets; define APIs and blockchain event schema.", "Digital Twin alpha and IPR agreement by M4."),
        ("M5-M8", "WP3 - AI optimisation and ledger module", "Develop forecasting models, sizing optimisation, predictive control and ledger-based event verification.", "Integrated prototype beta."),
        ("M9-M10", "WP4 - Testing and validation", "Test the solution with real-world data and multiple scenarios, including dynamic tariffs, high renewable generation and ancillary service activation.", "Test & Validation Report at M10."),
        ("M11-M14", "WP5 - Assessment and exploitation", "Assess technical, environmental and business impact; refine interface; prepare exploitation roadmap and commercial model.", "Business plan and exploitation roadmap at M14."),
    ]:
        cells = plan.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(plan, [1200, 2200, 4200, 1760])

    add_heading(doc, "8. Key Performance Indicators", 1)
    add_body(
        doc,
        "The KPIs are designed to be specific, measurable and time-bound across the 14-month programme, with baseline values confirmed during M1 and validation performed during M9-M10.",
    )
    kpis = doc.add_table(rows=1, cols=4)
    for idx, text in enumerate(["Dimension", "KPI", "Target", "Validation method"]):
        kpis.rows[0].cells[idx].text = text
    for row in [
        ("Resource optimisation", "Total electricity and heat cost", "10-15% reduction vs. baseline", "Digital Twin scenario comparison using real-world profiles."),
        ("Resource optimisation", "Storage CAPEX requirement", "8-12% reduction through optimal sizing", "Reference scenario vs. optimised configuration."),
        ("Green Deal", "Renewable self-consumption", "15-25% increase", "Measured simulated local use of renewable generation."),
        ("Green Deal", "Renewable curtailment", "10-20% reduction", "Scenario testing under high generation conditions."),
        ("Grid services", "Ancillary service fulfilment rate", ">90%", "Delivered service divided by requested service."),
        ("Grid services", "Activation success rate", ">95%", "Correct and timely response to activation events."),
        ("Trust and traceability", "Verified flexibility events", "100% recorded in ledger", "Audit of timestamped event records and hashes."),
        ("Exploitation", "Prototype readiness", "TRL 6-7 by M14", "Validated prototype and business roadmap."),
    ]:
        cells = kpis.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(kpis, [1600, 2700, 2200, 2860])

    add_heading(doc, "9. Data, testing and validation approach", 1)
    add_body(
        doc,
        "The solution will be tested in a digital environment using real-world electrical and thermal demand profiles, renewable generation profiles, asset price assumptions and dynamic tariff scenarios. The Digital Twin will compare baseline operation against optimised operation for multiple configurations, including different storage sizes, thermal storage strategies and grid service activation events.",
    )
    add_body(
        doc,
        "Validation will focus on reproducible simulation runs, transparent assumptions and quantified impact. The blockchain module will be validated by checking that each flexibility activation contains a timestamped request, asset response, delivered volume and verification hash, enabling independent audit without exposing commercially sensitive raw data.",
    )

    add_heading(doc, "10. Expected impact", 1)
    add_heading(doc, "10.1 Resource optimisation", 2)
    add_body(
        doc,
        "EnerTwin Ledger will reduce unnecessary storage oversizing, improve the timing of charge and discharge cycles, optimise thermal energy shifting and reduce operating cost under dynamic price signals. This supports more efficient local energy management and better utilisation of installed assets.",
    )
    add_heading(doc, "10.2 Green Deal impact", 2)
    add_body(
        doc,
        "The project supports higher renewable energy integration, lower curtailment, reduced fossil-based grid imports and more efficient coordination between electrical and thermal systems. These outcomes directly support decarbonisation, energy efficiency and resilience of local industrial energy systems.",
    )
    add_heading(doc, "10.3 Social and market impact", 2)
    add_body(
        doc,
        "The tool helps operators, SMEs and energy communities participate in flexibility markets with stronger transparency and lower technical barriers. It also supports workforce upskilling by translating complex microgrid operation into explainable digital twin scenarios and auditable decision records.",
    )

    add_heading(doc, "11. Business model and exploitation", 1)
    add_body(
        doc,
        "Target customers include industrial microgrid operators, renewable energy asset owners, energy communities, aggregators, ESCOs and grid service providers. The initial product will be commercialised as software-as-a-service with optional integration services. Pricing can combine a site subscription, setup fee and optional success fee linked to realised energy cost savings or flexibility market revenues.",
    )
    add_bullet(doc, "Primary market: industrial and commercial microgrids seeking cost reduction and renewable integration.")
    add_bullet(doc, "Secondary market: aggregators needing trusted flexibility verification across distributed assets.")
    add_bullet(doc, "Scale route: integration with existing energy management systems, SCADA gateways and AI-on-Demand ecosystem assets.")

    add_heading(doc, "12. Risks and mitigation", 1)
    risks = doc.add_table(rows=1, cols=3)
    for idx, text in enumerate(["Risk", "Impact", "Mitigation"]):
        risks.rows[0].cells[idx].text = text
    for row in [
        ("Limited data quality", "Forecasting and validation accuracy may decrease.", "Use data quality checks, fallback models and sensitivity analysis."),
        ("Complexity of thermal-electrical coupling", "Model may become hard to calibrate.", "Start with validated simplified models, then add detail incrementally."),
        ("Blockchain overhead", "Integration could become too heavy for the prototype.", "Use permissioned or lightweight ledger mode and store hashes, not large raw datasets."),
        ("Market access uncertainty", "Ancillary service rules differ by geography.", "Design modular verification records that can map to several market designs."),
    ]:
        cells = risks.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(risks, [2400, 3000, 3960])

    add_heading(doc, "13. Mandatory deliverables", 1)
    add_bullet(doc, "IPR Agreement established with playground owners by M4.")
    add_bullet(doc, "Test & Validation Report by M10.")
    add_bullet(doc, "Business plan and exploitation roadmap by M14.")

    add_heading(doc, "14. Sources used for call alignment", 1)
    add_body(doc, "AID4SME Open Call #2: https://aid4sme.eu/open-call-2/")
    add_body(doc, "AID4SME Annex 1.1 Challenge Description v1.1: https://aid4sme.eu/wp-content/uploads/2026/06/AID4SME_OC2_Annex-1.1_Challenge-description_v1.1.docx-1.pdf")
    add_body(doc, "EU CORDIS AID4SME fact sheet: https://cordis.europa.eu/project/id/101189562")

    props = doc.core_properties
    props.title = "EnerTwin Ledger - AID4SME Project Proposal"
    props.subject = "AID4SME Open Call #2 Challenge 2.3 proposal"
    props.author = ""
    props.comments = "Draft proposal prepared for AID4SME Open Call #2."
    props.keywords = "AID4SME, digital twin, blockchain, microgrid, energy optimisation, AI"

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
